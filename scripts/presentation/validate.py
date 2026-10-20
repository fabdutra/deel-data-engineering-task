"""Validate generated slides, source hashes and PowerPoint notes without Docker."""

import ast
import hashlib
import json
import datetime as dt
import statistics
from pathlib import Path
import fitz
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/presentation"


def main():
    manifest = json.loads((OUT / "manifest.json").read_text())
    for item in manifest["inputs"] + manifest["outputs"]:
        p = ROOT / item["path"]
        assert (
            p.exists() and hashlib.sha256(p.read_bytes()).hexdigest() == item["sha256"]
        ), item["path"]
    pdf = fitz.open(OUT / "acme-delivery-assessment.pdf")
    ppt = Presentation(OUT / "acme-delivery-assessment.pptx")
    assert len(pdf) == len(ppt.slides) == manifest["total_slides"] == 10
    assert manifest["main_slides"] == 10
    for i, (page, slide) in enumerate(zip(pdf, ppt.slides), 1):
        assert abs(page.rect.width - 960) < 1 and abs(page.rect.height - 540) < 1, i
        assert (
            page.get_text().strip() and slide.notes_slide.notes_text_frame.text.strip()
        ), i
        assert f"{i:02d} / 10" in page.get_text(), i
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (i, "negative shape position")
            assert (
                shape.left + shape.width <= ppt.slide_width + 100
                and shape.top + shape.height <= ppt.slide_height + 100
            ), (i, "shape exceeds slide")
        for block in page.get_text("blocks"):
            assert (
                block[0] >= -1
                and block[1] >= -1
                and block[2] <= 961
                and block[3] <= 541
            ), (i, "PDF text outside page")
    for folder in ("scripts/presentation", "scripts/stress"):
        for p in (ROOT / folder).glob("*.py"):
            ast.parse(p.read_text())
    proof = json.loads((ROOT / "docs/evidence/stress/final-integrity.json").read_text())
    a = proof["event_counters"]
    b = proof["physical_bronze_check"]
    assert (
        a["kafka_events"]
        == a["bronze"]
        == a["materialized"]
        == b["rows"]
        == b["unique_coordinates"]
        == 774001
    )
    assert proof["aggregate_reconciliation"]["passed"]
    # Recompute the analytical summaries from individual retained batch records.
    batches = json.loads(
        (ROOT / "docs/evidence/stress/streaming-batches.json").read_text()
    )
    levels = json.loads(
        (ROOT / "docs/evidence/stress/streaming-levels.json").read_text()
    )
    for level in levels:
        lo = dt.datetime.fromisoformat(level["started_at_utc"]).timestamp() * 1000
        hi = dt.datetime.fromisoformat(level["load_ended_at_utc"]).timestamp() * 1000
        selected = [
            b
            for b in batches
            if b.get("source_ts_min_ms") is not None
            and lo <= b["source_ts_min_ms"] <= hi
        ]
        ages = [b["source_to_publication_age_p95_seconds"] for b in selected]
        summary = level["freshness"]
        assert len(ages) == summary["batch_count"]
        assert (
            abs(statistics.median(ages) - summary["median_of_batch_p95_seconds"]) < 1e-6
        )
        assert abs(max(ages) - summary["maximum_batch_p95_seconds"]) < 1e-6
        assert (
            abs(
                level["emitted_events"] / level["elapsed_seconds"]
                - level["actual_events_per_second"]
            )
            < 1e-6
        )
    assert sum(b["input_events"] for b in batches) == 774001
    print(
        "Validated 10 PDF/PPTX slides, notes, geometry, hashes, Python syntax, integrity counts and recomputed freshness/rate summaries."
    )


if __name__ == "__main__":
    main()
