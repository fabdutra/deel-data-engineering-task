"""Build the assessment deck from repository evidence, using one layout for PDF/PPTX.

Run from any directory with the packages pinned in requirements.txt installed.
No network, Docker, credentials or private context files are used by this build.
"""

from pathlib import Path
import hashlib
import json
import math
import os
import sys

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/presentation"
os.environ.setdefault("MPLCONFIGDIR", str(OUT / "assets/.mpl-cache"))
import matplotlib

matplotlib.use("Agg")
from matplotlib import pyplot as plt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR

W, H = 960, 540
NAVY = "122B3B"
TEAL = "007F78"
MINT = "DCF3EC"
BG = "F6F8FA"
WHITE = "FFFFFF"
INK = "183342"
MUTED = "516978"
GOLD = "D49C3D"
LIGHT = "E5EBEF"
RED = "AD4D3D"
REPO = "https://github.com/fabdutra/deel-data-engineering-task"
FILES = set()


def load(path):
    FILES.add(path)
    return json.loads((ROOT / path).read_text())


def font_setup():
    candidates = [
        (
            "Arial",
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        ),
        ("Arial", "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        (
            "Liberation Sans",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        ),
        (
            "DejaVu Sans",
            str(Path(matplotlib.get_data_path()) / "fonts/ttf/DejaVuSans.ttf"),
            str(Path(matplotlib.get_data_path()) / "fonts/ttf/DejaVuSans-Bold.ttf"),
        ),
    ]
    for name, regular, bold in candidates:
        if Path(regular).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont("Deck", regular))
            pdfmetrics.registerFont(TTFont("DeckBold", bold))
            return name
    raise RuntimeError("No usable font found")


FONT = font_setup()
sys.modules.setdefault("build", sys.modules[__name__])


def wrap(value, width, size, bold=False):
    font = "DeckBold" if bold else "Deck"
    output = []
    for paragraph in str(value).split("\n"):
        line = ""
        for word in paragraph.split():
            trial = (line + " " + word).strip()
            if line and pdfmetrics.stringWidth(trial, font, size) > width:
                output.append(line)
                line = word
            else:
                line = trial
        output.append(line)
    return output


class Slide:
    def __init__(self, title, section="SOLUTION", sources=(), notes="", dark=False):
        self.title = title
        self.section = section
        self.sources = list(sources)
        self.notes = notes
        self.dark = dark
        self.items = []
        self.rect(0, 0, W, H, NAVY if dark else BG)
        if title:
            self.text(
                48, 27, 864, section.upper(), 10, TEAL if not dark else "6DE0C6", True
            )
            self.text(48, 53, 864, title, 28, WHITE if dark else NAVY, True)

    def rect(self, x, y, w, h, fill, stroke=None, radius=0):
        self.items.append(
            dict(
                kind="rect", x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, radius=radius
            )
        )

    def line(self, x1, y1, x2, y2, color=TEAL, width=1.8, arrow=False):
        self.items.append(
            dict(
                kind="line",
                x1=x1,
                y1=y1,
                x2=x2,
                y2=y2,
                color=color,
                width=width,
                arrow=arrow,
            )
        )

    def text(
        self, x, y, w, value, size=17, color=INK, bold=False, link=None, leading=None
    ):
        leading = leading or size * 1.28
        lines = wrap(value, w, size, bold)
        for i, line in enumerate(lines):
            self.items.append(
                dict(
                    kind="text",
                    x=x,
                    y=y + i * leading,
                    w=w,
                    h=leading + 4,
                    text=line,
                    size=size,
                    color=color,
                    bold=bold,
                    link=link,
                )
            )
        return len(lines) * leading

    def card(self, x, y, w, h, title, body="", fill=WHITE, accent=TEAL, size=16):
        self.rect(x, y, w, h, fill, radius=10)
        self.rect(x, y, 4, h, accent, radius=0)
        compact = h < 90
        used = self.text(
            x + 18,
            y + (12 if compact else 15),
            w - 34,
            title,
            15 if compact else 18,
            INK,
            True,
        )
        if body:
            self.text(
                x + 18, y + (18 if compact else 24) + used, w - 34, body, size, MUTED
            )

    def stat(self, x, y, w, value, label, sub="", color=TEAL):
        self.rect(x, y, w, 110, WHITE, radius=10)
        self.text(x + 16, y + 10, w - 30, value, 31, color, True)
        self.text(x + 16, y + 49, w - 30, label, 15, INK, True)
        if sub:
            self.text(x + 16, y + 74, w - 30, sub, 12, MUTED)

    def image(self, path, x, y, w, h):
        self.items.append(dict(kind="image", path=str(path), x=x, y=y, w=w, h=h))

    def table(self, x, y, widths, headers, rows, size=14, rowh=42):
        total = sum(widths)
        self.rect(x, y, total, 36, NAVY)
        cursor = x
        for width, value in zip(widths, headers):
            self.text(cursor + 10, y + 8, width - 20, value, size - 1, WHITE, True)
            cursor += width
        top = y + 36
        for i, row in enumerate(rows):
            rh = max(
                rowh,
                max(len(wrap(v, w - 20, size)) for v, w in zip(row, widths))
                * size
                * 1.28
                + 18,
            )
            self.rect(x, top, total, rh, WHITE if i % 2 == 0 else "EAF0F3")
            cursor = x
            for width, value in zip(widths, row):
                self.text(cursor + 10, top + 8, width - 20, value, size, INK)
                cursor += width
            top += rh
        return top

    def finish(self, index, total):
        color = "B2CAD5" if self.dark else MUTED
        self.line(48, 498, 912, 498, "375361" if self.dark else LIGHT, 0.7)
        source = (
            "Sources: " + "  |  ".join(self.sources)
            if self.sources
            else "Fabricio Dutra  •  Deel Senior Data Engineer assessment"
        )
        self.text(
            48,
            508,
            792,
            source,
            9.1,
            color,
            link=REPO + "/blob/main/" + self.sources[0] if self.sources else REPO,
        )
        self.text(865, 508, 48, f"{index:02d} / {total:02d}", 10, color, True)


def charts(data):
    assets = OUT / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 11,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.spines.left": False,
            "axes.edgecolor": "#CDD8DE",
            "axes.labelcolor": "#516978",
            "xtick.color": "#516978",
            "ytick.color": "#516978",
            "figure.facecolor": "#FFFFFF",
            "axes.facecolor": "#FFFFFF",
        }
    )
    uniform = sorted(
        [r for r in data["materialization"] if not r["hot"]], key=lambda r: r["events"]
    )
    fig, ax = plt.subplots(figsize=(7.1, 3.75))
    fig.subplots_adjust(left=0.09, right=0.98, top=0.93, bottom=0.18)
    silver = [r["phases"][0]["silver_merge_seconds"] for r in uniform]
    gold = [r["phases"][0]["gold_seconds"] for r in uniform]
    other = [r["initial_seconds"] - a - b for r, a, b in zip(uniform, silver, gold)]
    xs = range(3)
    ax.bar(xs, silver, color="#8BBCAF", label="Silver updates", width=0.52)
    ax.bar(xs, gold, bottom=silver, color="#007F78", label="Gold tables", width=0.52)
    ax.bar(
        xs,
        other,
        bottom=[a + b for a, b in zip(silver, gold)],
        color="#CBD6DE",
        label="Validation / other work",
        width=0.52,
    )
    for x, r in zip(xs, uniform):
        ax.text(
            x,
            r["initial_seconds"] + 6,
            f"{r['initial_seconds']:.1f}s",
            ha="center",
            weight="bold",
            color="#183342",
            fontsize=12,
        )
    ax.set_xticks(list(xs), ["0.6m events", "3.0m events", "6.0m events"])
    ax.set_ylabel("Initial processing time (seconds)")
    ax.set_ylim(0, 420)
    ax.grid(axis="y", alpha=0.15)
    ax.set_axisbelow(True)
    ax.tick_params(axis="both", length=0)
    ax.legend(loc="upper left", frameon=False, fontsize=10)
    fig.savefig(assets / "materialization.png", dpi=220)
    plt.close(fig)
    levels = data["levels"]
    fig, ax = plt.subplots(figsize=(7.1, 3.75))
    fig.subplots_adjust(left=0.09, right=0.98, top=0.91, bottom=0.2)
    median = [r["freshness"]["median_of_batch_p95_seconds"] for r in levels]
    maximum = [r["freshness"]["maximum_batch_p95_seconds"] for r in levels]
    xs = [r["actual_events_per_second"] for r in levels]
    ax.plot(
        xs,
        median,
        "o-",
        color="#007F78",
        linewidth=2.8,
        label="Batch p95: median",
    )
    ax.plot(
        xs,
        maximum,
        "o--",
        color="#D49C3D",
        linewidth=2,
        label="Batch p95: highest",
    )
    ax.axhline(
        30, color="#AD4D3D", ls=":", linewidth=1.7, label="30s test reference"
    )
    for x, m in zip(xs, maximum):
        ax.text(x, m + 4, f"{m:.1f}s", ha="center", color="#183342", fontsize=10)
    ax.set_xticks(list(xs), [f"{r['actual_events_per_second']:,.0f}" for r in levels])
    ax.set_xlabel("Actual source events / second")
    ax.set_ylabel("Time until published (seconds)")
    ax.set_ylim(0, 140)
    ax.legend(loc="upper left", frameon=False, fontsize=9)
    ax.grid(axis="y", alpha=0.15)
    ax.tick_params(axis="both", length=0)
    ax.set_axisbelow(True)
    fig.savefig(assets / "freshness.png", dpi=220)
    plt.close(fig)
    return assets


def render(slides):
    prs = Presentation()
    prs.slide_width = Pt(W)
    prs.slide_height = Pt(H)
    prs.core_properties.title = "ACME Delivery: streaming operational analytics"
    prs.core_properties.subject = "Deel Senior Data Engineer assessment"
    prs.core_properties.author = "Fabricio Dutra"
    pdf = canvas.Canvas(
        str(OUT / "acme-delivery-assessment.pdf"), pagesize=(W, H), pageCompression=1
    )
    pdf.setTitle(prs.core_properties.title)
    pdf.setAuthor("Fabricio Dutra")
    pdf.setSubject("Architecture, methodology, Spark sizing and measured results")
    issues = []
    for number, slide in enumerate(slides, 1):
        slide.finish(number, len(slides))
        ps = prs.slides.add_slide(prs.slide_layouts[6])
        for item in slide.items:
            kind = item["kind"]
            if kind == "rect":
                x, y, w, h = (item[k] for k in ("x", "y", "w", "h"))
                pdf.setFillColor(HexColor("#" + item["fill"]))
                pdf.setStrokeColor(HexColor("#" + (item["stroke"] or item["fill"])))
                pdf.roundRect(
                    x,
                    H - y - h,
                    w,
                    h,
                    item["radius"],
                    fill=1,
                    stroke=bool(item["stroke"]),
                )
                shape = ps.shapes.add_shape(
                    (
                        MSO_SHAPE.ROUNDED_RECTANGLE
                        if item["radius"]
                        else MSO_SHAPE.RECTANGLE
                    ),
                    Pt(x),
                    Pt(y),
                    Pt(w),
                    Pt(h),
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(item["fill"])
                if item["stroke"]:
                    shape.line.color.rgb = RGBColor.from_string(item["stroke"])
                else:
                    shape.line.fill.background()
                if item["radius"]:
                    shape.adjustments[0] = 0.12
            elif kind == "text":
                x, y, w, h = (item[k] for k in ("x", "y", "w", "h"))
                size = item["size"]
                pdf.setFont("DeckBold" if item["bold"] else "Deck", size)
                pdf.setFillColor(HexColor("#" + item["color"]))
                pdf.drawString(x, H - y - size, item["text"])
                if item.get("link"):
                    pdf.linkURL(
                        item["link"],
                        (x, H - y - h, x + w, H - y),
                        relative=0,
                        thickness=0,
                    )
                box = ps.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
                tf = box.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                tf.vertical_anchor = MSO_ANCHOR.TOP
                para = tf.paragraphs[0]
                para.space_before = para.space_after = Pt(0)
                run = para.add_run()
                run.text = item["text"]
                run.font.name = FONT
                run.font.size = Pt(size)
                run.font.bold = item["bold"]
                run.font.color.rgb = RGBColor.from_string(item["color"])
                if item.get("link"):
                    run.hyperlink.address = item["link"]
                if x < 0 or y < 0 or x + w > W + 0.5 or y + h > H + 0.5:
                    issues.append((number, item["text"], "out of bounds"))
                if (
                    pdfmetrics.stringWidth(
                        item["text"], "DeckBold" if item["bold"] else "Deck", size
                    )
                    > w + 0.5
                ):
                    issues.append((number, item["text"], "line too wide"))
            elif kind == "line":
                x1, y1, x2, y2 = (item[k] for k in ("x1", "y1", "x2", "y2"))
                color = item["color"]
                width = item["width"]
                pdf.setStrokeColor(HexColor("#" + color))
                pdf.setLineWidth(width)
                pdf.line(x1, H - y1, x2, H - y2)
                shape = ps.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)
                )
                shape.line.color.rgb = RGBColor.from_string(color)
                shape.line.width = Pt(width)
                if item["arrow"]:
                    angle = math.atan2(y2 - y1, x2 - x1)
                    points = [
                        (x2, y2),
                        (
                            x2 - 7 * math.cos(angle) + 3 * math.sin(angle),
                            y2 - 7 * math.sin(angle) - 3 * math.cos(angle),
                        ),
                        (
                            x2 - 7 * math.cos(angle) - 3 * math.sin(angle),
                            y2 - 7 * math.sin(angle) + 3 * math.cos(angle),
                        ),
                    ]
                    pdf.setFillColor(HexColor("#" + color))
                    p = pdf.beginPath()
                    p.moveTo(points[0][0], H - points[0][1])
                    p.lineTo(points[1][0], H - points[1][1])
                    p.lineTo(points[2][0], H - points[2][1])
                    p.close()
                    pdf.drawPath(p, fill=1, stroke=0)
                    tri = ps.shapes.add_shape(
                        MSO_SHAPE.ISOSCELES_TRIANGLE,
                        Pt(x2 - 4),
                        Pt(y2 - 4),
                        Pt(8),
                        Pt(8),
                    )
                    tri.rotation = math.degrees(angle) + 90
                    tri.fill.solid()
                    tri.fill.fore_color.rgb = RGBColor.from_string(color)
                    tri.line.fill.background()
            elif kind == "image":
                pdf.drawImage(
                    item["path"],
                    item["x"],
                    H - item["y"] - item["h"],
                    item["w"],
                    item["h"],
                    mask="auto",
                )
                ps.shapes.add_picture(
                    item["path"],
                    Pt(item["x"]),
                    Pt(item["y"]),
                    width=Pt(item["w"]),
                    height=Pt(item["h"]),
                )
        ps.notes_slide.notes_text_frame.text = (
            slide.notes + "\n\nEvidence: " + ", ".join(slide.sources)
        )
        pdf.showPage()
    if issues:
        raise RuntimeError("Layout validation failed: " + repr(issues))
    pdf.save()
    prs.save(OUT / "acme-delivery-assessment.pptx")
    return prs


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    data = {
        "materialization": load("docs/evidence/stress/materialization.json"),
        "levels": load("docs/evidence/stress/streaming-levels.json"),
        "integrity": load("docs/evidence/stress/final-integrity.json"),
        "query": load("docs/evidence/stress/query-output.json"),
        "recovery": load("docs/evidence/stress/broker-outage.jsonl"),
    }
    from deck import make_slides

    assets = charts(data)
    slides = make_slides(Slide, data, assets)
    render(slides)
    for slide in slides:
        FILES.update(slide.sources)
    FILES.update(
        str(p.relative_to(ROOT))
        for p in (ROOT / "scripts/stress").iterdir()
        if p.is_file() and p.suffix in (".py", ".sh", ".yaml", ".md")
    )
    FILES.update(
        str(p.relative_to(ROOT))
        for p in (ROOT / "docs/evidence/stress").iterdir()
        if p.is_file()
    )
    FILES.update(
        str(p.relative_to(ROOT))
        for p in (ROOT / "scripts/presentation").iterdir()
        if p.is_file() and p.suffix in (".py", ".txt")
    )
    files = sorted(FILES)
    manifest = {
        "title": "ACME Delivery: streaming operational analytics",
        "author": "Fabricio Dutra",
        "language": "en",
        "main_slides": len(slides),
        "total_slides": len(slides),
        "format": "16:9",
        "measurement_dates": "2026-10-19 to 2026-10-20; evidence timestamps UTC",
        "baseline_evidence_commit": "00b63a274df87feaf0991b2fc7e3674dadea852e",
        "font": FONT,
        "slides": [
            {
                "number": i,
                "title": s.title or "ACME Delivery: streaming operational analytics",
                "sources": s.sources,
            }
            for i, s in enumerate(slides, 1)
        ],
        "inputs": [
            {"path": p, "sha256": hashlib.sha256((ROOT / p).read_bytes()).hexdigest()}
            for p in files
        ],
        "outputs": [
            {
                "path": str(p.relative_to(ROOT)),
                "sha256": hashlib.sha256(p.read_bytes()).hexdigest(),
            }
            for p in (
                OUT / "acme-delivery-assessment.pdf",
                OUT / "acme-delivery-assessment.pptx",
            )
        ],
    }
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    (OUT / "slides.md").write_text(
        "# Presentation outline and English speaker notes\n\n"
        + "\n\n".join(
            f"## {i:02d}. {s.title or 'ACME Delivery: streaming operational analytics'}\n\n{s.notes}\n\nSources: "
            + ", ".join(f"[{p}]({REPO}/blob/main/{p})" for p in s.sources)
            for i, s in enumerate(slides, 1)
        )
        + "\n"
    )
    print(
        f"Generated {len(slides)} slides in PDF and PowerPoint, with source manifest and speaker notes."
    )


if __name__ == "__main__":
    main()
